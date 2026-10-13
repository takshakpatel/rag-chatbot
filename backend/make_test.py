import sys
from pptx import Presentation
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "Hello World Presentation"
slide.placeholders[1].text = "This is a test slide content for Lumen."
prs.save("E:/Project/backend/test.pptx")
print("PPTX created successfully")
