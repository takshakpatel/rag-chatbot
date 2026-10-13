from __future__ import annotations

import os
from pathlib import Path

from langchain_community.document_loaders import (
    PyPDFLoader,
    Docx2txtLoader,
    BSHTMLLoader,
    TextLoader,
)
from langchain_core.documents import Document

from app.image_extractor import extract_images_from_pdf


def _load_pptx(path: Path) -> list[Document]:
    """Load a PowerPoint file, extracting text from every slide."""
    try:
        from pptx import Presentation  # type: ignore
    except ImportError:
        print("Warning: python-pptx not installed. Run: pip install python-pptx")
        return []

    try:
        prs = Presentation(str(path))
        docs = []
        for slide_num, slide in enumerate(prs.slides, start=1):
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
            if texts:
                content = "\n".join(texts)
                docs.append(Document(
                    page_content=content,
                    metadata={
                        "source": str(path),
                        "filename": path.name,
                        "slide": slide_num,
                    },
                ))
        return docs
    except Exception as e:
        print(f"Warning: failed to load {path}: {e}")
        return []


SUPPORTED_EXTENSIONS = {
    ".pdf": PyPDFLoader,
    ".docx": Docx2txtLoader,
    ".html": BSHTMLLoader,
    ".htm": BSHTMLLoader,
    ".md": TextLoader,
    ".txt": TextLoader,
    ".pptx": None,  # handled by _load_pptx
}


def load_documents(path: str | Path, extract_images: bool = True) -> list[Document]:
    """Load documents from a file or directory.

    Supported formats: PDF, DOCX, HTML, Markdown, TXT, PPTX.
    For PDFs, also extracts images (with captions) as separate Document objects.
    Returns a flat list of LangChain Document objects.
    """
    target = Path(path)
    docs: list[Document] = []

    if target.is_file():
        docs.extend(_load_file(target))
        if extract_images and target.suffix.lower() == ".pdf":
            docs.extend(extract_images_from_pdf(target))
    elif target.is_dir():
        for f in sorted(target.rglob("*")):
            if f.suffix.lower() in SUPPORTED_EXTENSIONS:
                docs.extend(_load_file(f))
                if extract_images and f.suffix.lower() == ".pdf":
                    docs.extend(extract_images_from_pdf(f))
    else:
        raise FileNotFoundError(f"Path does not exist: {target}")

    return docs


def load_uploaded_files(paths: list[str | Path], extract_images: bool = True) -> list[Document]:
    """Load documents from a list of individual files (e.g. user uploads).

    Each path must be a file with a supported extension. Unsupported files are
    skipped with a warning; returns a flat list of Document objects.
    """
    docs: list[Document] = []
    for raw in paths:
        p = Path(raw)
        if not p.is_file():
            print(f"Warning: skipped non-file: {p}")
            continue
        if p.suffix.lower() not in SUPPORTED_EXTENSIONS:
            print(f"Warning: unsupported file type, skipped: {p.name}")
            continue
        docs.extend(_load_file(p))
        if extract_images and p.suffix.lower() == ".pdf":
            docs.extend(extract_images_from_pdf(p))
    return docs


def _load_file(path: Path) -> list[Document]:
    ext = path.suffix.lower()

    if ext == ".pptx":
        return _load_pptx(path)

    loader_cls = SUPPORTED_EXTENSIONS.get(ext)
    if loader_cls is None:
        return []

    try:
        loader = loader_cls(str(path))
        docs = loader.load()
        for doc in docs:
            doc.metadata.setdefault("source", str(path))
            doc.metadata.setdefault("filename", path.name)
        return docs
    except Exception as e:
        print(f"Warning: failed to load {path}: {e}")
        return []
